"""Relatórios em PDF e Excel a partir de um resumo e de tabelas simples.

Cada tabela é {"title": str, "columns": [(nome, tipo)], "rows": [[...]]}, com tipo
"text", "int", "money" (centavos) ou "date" (ISO). O resumo é uma lista de (rótulo, valor, tipo).
"""
import io
from datetime import datetime
from zoneinfo import ZoneInfo

ZONE = ZoneInfo("America/Sao_Paulo")
TEAL, ORANGE, LIGHT = (16, 63, 66), (255, 122, 26), (244, 246, 241)


def brl(cents):
    value = f"{abs(cents) / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-R$ " if cents < 0 else "R$ ") + value


def show(value, kind):
    if value is None or value == "":
        return "-" if kind != "text" else ""
    if kind == "money":
        return brl(value)
    if kind == "date":
        return datetime.fromisoformat(value).astimezone(ZONE).strftime("%d/%m/%Y %H:%M")
    return str(value)


def latin(text):
    """As fontes padrão do PDF usam Latin-1; troca símbolos que ficariam como '?'."""
    for a, b in (("•", "-"), ("—", "-"), ("–", "-"), ("−", "-"), ("“", '"'), ("”", '"'), ("’", "'"), ("…", "...")):
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


def pdf(title, subtitle, summary, tables):
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    class Report(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", size=8)
            self.set_text_color(120, 130, 130)
            self.cell(0, 6, latin(f"SeuComércioAqui • {title} • página {self.page_no()}/{{nb}}"), align="C")

    doc = Report(orientation="P", unit="mm", format="A4")
    doc.set_auto_page_break(True, margin=16)
    doc.set_title(latin(title))
    doc.add_page()
    doc.set_fill_color(*TEAL)
    doc.rect(0, 0, 210, 30, "F")
    doc.set_fill_color(*ORANGE)
    doc.rect(0, 30, 210, 1.6, "F")
    doc.set_xy(12, 8)
    doc.set_text_color(255, 255, 255)
    doc.set_font("Helvetica", "B", 17)
    doc.cell(0, 8, latin(title))
    doc.set_xy(12, 17)
    doc.set_font("Helvetica", size=10)
    generated = datetime.now(ZONE).strftime("%d/%m/%Y %H:%M")
    doc.cell(0, 6, latin(f"{subtitle} • gerado em {generated}"))
    doc.set_y(38)
    doc.set_text_color(20, 44, 48)

    if summary:
        width = (210 - 24 - 9) / 4
        for index, (label, value, kind) in enumerate(summary):
            column, line = index % 4, index // 4
            x, y = 12 + column * (width + 3), 38 + line * 20
            doc.set_fill_color(*LIGHT)
            doc.rect(x, y, width, 17, "F")
            doc.set_xy(x + 3, y + 2.5)
            doc.set_font("Helvetica", size=7.5)
            doc.set_text_color(80, 104, 107)
            doc.cell(width - 6, 4, latin(label))
            doc.set_xy(x + 3, y + 8)
            doc.set_font("Helvetica", "B", 12)
            doc.set_text_color(20, 44, 48)
            doc.cell(width - 6, 6, latin(show(value, kind)))
        doc.set_y(38 + ((len(summary) + 3) // 4) * 20 + 2)

    for table in tables:
        if doc.get_y() > 250:
            doc.add_page()
        doc.ln(3)
        doc.set_font("Helvetica", "B", 12)
        doc.set_text_color(*TEAL)
        doc.cell(0, 7, latin(table["title"]), new_x="LMARGIN", new_y="NEXT")
        doc.set_text_color(20, 44, 48)
        if not table["rows"]:
            doc.set_font("Helvetica", size=9)
            doc.cell(0, 6, latin("Nada no período."), new_x="LMARGIN", new_y="NEXT")
            continue
        doc.set_font("Helvetica", size=8)
        weights = [{"text": 3, "date": 3, "money": 2.2}.get(kind, 1.4) for _, kind in table["columns"]]
        aligns = ["LEFT" if kind in ("text", "date") else "RIGHT" for _, kind in table["columns"]]
        with doc.table(col_widths=weights, text_align=aligns, line_height=5.2, padding=1.2,
                       headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=TEAL),
                       cell_fill_color=LIGHT, cell_fill_mode="ROWS", borders_layout="NONE", first_row_as_headings=True) as grid:
            header = grid.row()
            for name, _ in table["columns"]:
                header.cell(latin(name))
            for values in table["rows"]:
                line = grid.row()
                for value, (_, kind) in zip(values, table["columns"]):
                    line.cell(latin(show(value, kind)))
    doc.alias_nb_pages()
    return bytes(doc.output())


def xlsx(title, subtitle, summary, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    teal = PatternFill("solid", fgColor="103F42")
    money_format = '"R$" #,##0.00;-"R$" #,##0.00'
    book = Workbook()
    sheet = book.active
    sheet.title = "Resumo"
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=16, color="103F42")
    sheet["A2"] = f"{subtitle} • gerado em {datetime.now(ZONE).strftime('%d/%m/%Y %H:%M')}"
    sheet["A2"].font = Font(color="52686B")
    for row, (label, value, kind) in enumerate(summary, start=4):
        sheet.cell(row, 1, label).font = Font(bold=True)
        cell = sheet.cell(row, 2, value / 100 if kind == "money" and value is not None else value)
        if kind == "money":
            cell.number_format = money_format
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 20

    used = set()
    for table in tables:
        name = "".join(c for c in table["title"] if c not in '[]:*?/\\')[:31] or "Tabela"
        while name in used:
            name = name[:28] + str(len(used))
        used.add(name)
        ws = book.create_sheet(name)
        for col, (header, _) in enumerate(table["columns"], start=1):
            cell = ws.cell(1, col, header)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), teal
            cell.alignment = Alignment(vertical="center")
        for row, values in enumerate(table["rows"], start=2):
            for col, (value, (_, kind)) in enumerate(zip(values, table["columns"]), start=1):
                if kind == "money" and value is not None:
                    cell = ws.cell(row, col, value / 100)
                    cell.number_format = money_format
                elif kind == "date" and value:
                    cell = ws.cell(row, col, datetime.fromisoformat(value).astimezone(ZONE).replace(tzinfo=None))
                    cell.number_format = "dd/mm/yyyy hh:mm"
                else:
                    ws.cell(row, col, value)
        for col, (header, kind) in enumerate(table["columns"], start=1):
            longest = max([len(header)] + [len(show(r[col - 1], kind)) for r in table["rows"][:300]])
            ws.column_dimensions[get_column_letter(col)].width = min(48, max(10, longest + 2))
        ws.freeze_panes = "A2"
        if table["rows"]:
            ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def build(fmt, title, subtitle, summary, tables):
    if fmt == "pdf":
        return pdf(title, subtitle, summary, tables), "application/pdf", "pdf"
    return (xlsx(title, subtitle, summary, tables),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx")
