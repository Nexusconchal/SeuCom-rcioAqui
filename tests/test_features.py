import uuid
from datetime import datetime

from app import ZONE, store_status
from tests.test_api import app, shop, send, register, payload
from tests.test_business import make_order, advance


def schedule_store(hours, open_=1, enabled=1):
    import json
    return {"enabled": enabled, "open": open_, "auto_hours": 1, "hours": json.dumps(hours)}


def at(weekday_day, hour, minute=0):
    # 2026-10-05 é uma segunda-feira.
    return datetime(2026, 10, 5 + weekday_day, hour, minute, tzinfo=ZONE)


def test_store_status_follows_schedule_and_overnight_shifts():
    store = schedule_store({"0": [["11:00", "14:00"], ["18:00", "02:00"]], "4": [["18:00", "23:00"]]})
    assert store_status(store, at(0, 12)) == (True, "Aberto até 14:00")
    assert store_status(store, at(0, 15)) == (False, "Fechado • abre hoje às 18:00")
    assert store_status(store, at(0, 23, 30))[0] is True
    assert store_status(store, at(1, 1, 30)) == (True, "Aberto até 02:00")
    assert store_status(store, at(1, 3)) == (False, "Fechado • abre sexta às 18:00")
    assert store_status(schedule_store({"0": [["11:00", "14:00"]]}, open_=0), at(0, 12))[0] is False
    assert store_status(schedule_store({"0": [["11:00", "14:00"]]}, enabled=0), at(0, 12))[0] is False


def test_hours_validation_and_closed_schedule_blocks_orders(shop):
    client, pid = shop
    store = client.get("/api/admin/store").json["store"]
    assert send(client, "PUT", "/api/admin/store", {**store, "auto_hours": True, "hours": {}}).status_code == 400
    assert send(client, "PUT", "/api/admin/store", {**store, "hours": {"0": [["25:00", "26:00"]]}}).status_code == 400
    assert send(client, "PUT", "/api/admin/store", {**store, "hours": {"9": [["10:00", "11:00"]]}}).status_code == 400
    weekday = str((datetime.now(ZONE).weekday() + 3) % 7)
    saved = send(client, "PUT", "/api/admin/store", {**store, "auto_hours": True, "hours": {weekday: [["10:00", "11:00"]]}})
    assert saved.status_code == 200 and saved.json["store"]["open_now"] is False
    assert "abre" in saved.json["store"]["status_text"]
    public = client.get("/api/store/minha-loja").json["store"]
    assert public["open_now"] is False and "blocked_reason" not in public
    assert send(client, "POST", "/api/store/minha-loja/orders", payload(pid), uuid.uuid4().hex).status_code == 409
    # Pedido de balcão continua permitido para a própria loja.
    assert send(client, "POST", "/api/admin/orders", payload(pid), uuid.uuid4().hex).status_code == 201


def test_delivery_fee_by_neighborhood(shop):
    client, pid = shop
    store = client.get("/api/admin/store").json["store"]
    zones = [{"name": "Centro", "fee": 300}, {"name": "Jardim Bela Vista", "fee": 800}]
    assert send(client, "PUT", "/api/admin/store", {**store, "delivery_zones": zones + [{"name": "centro", "fee": 1}]}).status_code == 400
    assert send(client, "PUT", "/api/admin/store", {**store, "delivery_zones": zones}).status_code == 200
    assert send(client, "POST", "/api/store/minha-loja/quote", payload(pid)).status_code == 400
    quote = send(client, "POST", "/api/store/minha-loja/quote", payload(pid, zone="Jardim Bela Vista")).json
    assert quote["delivery_fee"] == 800 and quote["zone"] == "Jardim Bela Vista"
    pickup = send(client, "POST", "/api/store/minha-loja/quote", payload(pid, mode="pickup")).json
    assert pickup["delivery_fee"] == 0
    order = make_order(client, pid, zone="Centro")
    assert order["delivery_fee"] == 300 and order["zone"] == "Centro"


def test_store_link_preview_and_qrcode(shop):
    client, _ = shop
    page = client.get("/loja/minha-loja")
    html = page.get_data(as_text=True)
    assert page.status_code == 200 and "<title>Minha Loja • Peça pelo cardápio digital</title>" in html
    assert 'property="og:image"' in html and "/static/og-image.png" in html
    assert client.get("/loja/nao-existe").status_code == 404
    svg = client.get("/api/admin/qrcode")
    assert svg.status_code == 200 and svg.mimetype == "image/svg+xml" and b"<svg" in svg.data
    png = client.get("/api/admin/qrcode?format=png&download=1")
    assert png.data[:8] == b"\x89PNG\r\n\x1a\n" and "attachment" in png.headers["Content-Disposition"]
    assert app_client_without_login(client).get("/api/admin/qrcode").status_code == 401


def app_client_without_login(client):
    return client.application.test_client()


def test_platform_monthly_block_promo_and_notices(app, shop):
    client, pid = shop
    other = app.test_client()
    assert register(other, "outra-loja", "other@example.com").status_code == 201
    assert other.get("/api/platform/monthly").status_code == 403
    app.test_cli_runner().invoke(args=["grant-platform-admin", "--email", "owner@example.com"])
    sid = client.get("/api/admin/store").json["store"]["id"]
    order = make_order(client, pid)
    advance(client, order["id"], "preparing", "ready", "delivering", "completed")

    report = client.get("/api/platform/monthly?months=3").json
    assert len(report["months"]) == 3 and report["month"] == datetime.now(ZONE).strftime("%Y-%m")
    current = report["months"][-1]
    assert current["orders"] == 1 and current["gmv"] == order["total"] and current["new_stores"] == 2
    assert report["ranking"][0]["id"] == sid and report["ranking"][0]["gmv"] == order["total"]
    assert client.get("/api/platform/monthly?month=2026-13").status_code == 400

    # Promoção na mensalidade e inadimplência do mês.
    contract = {"commission_bps": 0, "monthly_fee": 9900, "enabled": True}
    assert send(client, "PUT", f"/api/platform/stores/{sid}", {**contract, "promo_fee": 0}).status_code == 400
    assert send(client, "PUT", f"/api/platform/stores/{sid}", {**contract, "promo_fee": 4900, "promo_until": "2999-12-31", "promo_label": "Lançamento"}).status_code == 200
    summary = client.get("/api/platform/summary").json
    mine = next(s for s in summary["stores"] if s["id"] == sid)
    assert mine["effective_fee"] == 4900 and mine["promo_active"] and mine["overdue"] and summary["overdue"] == 4900
    assert send(client, "POST", "/api/platform/ledger", {"store_id": sid, "kind": "receipt", "description": "Mensalidade", "amount": 4900}, uuid.uuid4().hex).status_code == 201
    assert next(s for s in client.get("/api/platform/summary").json["stores"] if s["id"] == sid)["overdue"] is False

    # Bloqueio com motivo e desbloqueio.
    other_id = other.get("/api/admin/store").json["store"]["id"]
    assert send(client, "POST", f"/api/platform/stores/{other_id}/block", {"blocked": True, "reason": ""}).status_code == 400
    assert send(client, "POST", f"/api/platform/stores/{other_id}/block", {"blocked": True, "reason": "Mensalidade atrasada"}).status_code == 200
    assert other.get("/api/admin/notices").json["blocked_reason"] == "Mensalidade atrasada"
    settings = other.get("/api/admin/store").json["store"]
    assert send(other, "PUT", "/api/admin/store", {**settings, "open": True}).status_code == 403
    assert send(client, "POST", f"/api/platform/stores/{other_id}/block", {"blocked": False, "reason": ""}).status_code == 200
    assert other.get("/api/admin/notices").json["blocked_reason"] == ""

    # Avisos e promoções para todos os lojistas.
    assert send(other, "POST", "/api/platform/notices", {"title": "Oi", "body": "Teste", "kind": "info"}).status_code == 403
    created = send(client, "POST", "/api/platform/notices", {"title": "Indique e ganhe", "body": "Indique uma loja e ganhe 1 mês grátis.", "kind": "promo", "expires": "2999-01-01"})
    assert created.status_code == 201
    assert send(client, "POST", "/api/platform/notices", {"title": "Antigo", "body": "Já venceu.", "kind": "info", "expires": "2000-01-01"}).status_code == 201
    notices = other.get("/api/admin/notices").json["notices"]
    assert [n["title"] for n in notices] == ["Indique e ganhe"]
    assert send(client, "PATCH", f"/api/platform/notices/{created.json['id']}", {"active": False}).status_code == 200
    assert other.get("/api/admin/notices").json["notices"] == []


def test_owner_site_is_separate_and_not_indexed(app, shop):
    client, _ = shop
    page = client.get("/admin")
    assert page.status_code == 200 and "noindex" in page.headers["X-Robots-Tag"]
    old = client.get("/plataforma")
    assert old.status_code == 301 and old.headers["Location"].endswith("/admin")
    assert client.get("/api/platform/report").status_code == 403


def test_merchant_and_platform_reports_in_pdf_and_excel(app, shop):
    import io
    from openpyxl import load_workbook
    client, pid = shop
    order = make_order(client, pid)
    advance(client, order["id"], "preparing", "ready", "delivering", "completed")
    send(client, "POST", "/api/admin/expenses", {"description": "Gás de cozinha", "amount": 12000})
    pdf = client.get("/api/admin/report?days=30&format=pdf")
    assert pdf.status_code == 200 and pdf.data[:5] == b"%PDF-" and pdf.mimetype == "application/pdf"
    assert "relatorio-minha-loja-30-dias.pdf" in pdf.headers["Content-Disposition"]
    excel = client.get("/api/admin/report?month=" + datetime.now(ZONE).strftime("%Y-%m") + "&format=xlsx")
    book = load_workbook(io.BytesIO(excel.data))
    assert {"Resumo", "Pedidos", "Despesas", "Produtos mais vendidos"} <= set(book.sheetnames)
    assert book["Pedidos"].max_row == 2 and book["Pedidos"]["I2"].value == order["total"] / 100
    assert client.get("/api/admin/report?days=30&format=doc").status_code == 400
    app.test_cli_runner().invoke(args=["grant-platform-admin", "--email", "owner@example.com"])
    platform_pdf = client.get("/api/platform/report?format=pdf")
    assert platform_pdf.status_code == 200 and platform_pdf.data[:5] == b"%PDF-"
    platform_xlsx = load_workbook(io.BytesIO(client.get("/api/platform/report?format=xlsx").data))
    assert platform_xlsx["Últimos 12 meses"].max_row == 13
    assert client.get("/api/platform/report?month=2001-01").status_code == 400


def test_api_keys_rename_rotate_test_and_delete(app, shop):
    client, pid = shop
    created = send(client, "POST", "/api/admin/integrations/keys", {"name": "MotoJá", "scopes": ["orders:read", "deliveries:write"]}).json
    kid, token = created["id"], created["token"]
    assert send(client, "POST", "/api/admin/integrations/test", {"token": token}).json["ok"] is True
    assert send(client, "POST", "/api/admin/integrations/test", {"token": "abc"}).json["ok"] is False
    assert send(client, "PATCH", f"/api/admin/integrations/keys/{kid}", {"name": "MotoJá Conchal"}).status_code == 200
    rotated = send(client, "POST", f"/api/admin/integrations/keys/{kid}/rotate", {}).json["token"]
    partner = app.test_client()
    assert partner.get("/api/v1/store", headers={"Authorization": "Bearer " + token}).status_code == 401
    assert partner.get("/api/v1/store", headers={"Authorization": "Bearer " + rotated}).status_code == 200
    # Uma chave já usada em entrega sai da lista, mas o vínculo do pedido continua.
    order = make_order(client, pid)
    advance(client, order["id"], "preparing", "ready")
    auth = {"Authorization": "Bearer " + rotated}
    assert partner.post(f"/api/v1/deliveries/{order['id']}/claim", json={"external_id": "c-1", "provider": "motoja"}, headers=auth).status_code == 201
    listed = client.get("/api/admin/integrations").json
    assert listed["keys"][0]["name"] == "MotoJá Conchal" and listed["keys"][0]["deliveries"] == 1
    assert send(client, "DELETE", f"/api/admin/integrations/keys/{kid}/permanent", {}).status_code == 200
    after = client.get("/api/admin/integrations").json
    assert after["keys"] == [] and len(after["deliveries"]) == 1
    assert partner.get("/api/v1/store", headers=auth).status_code == 401
    assert send(client, "POST", "/api/admin/integrations/test", {"token": rotated}).json["ok"] is False
    # Chave nunca usada é apagada de vez; outra loja não consegue mexer.
    unused = send(client, "POST", "/api/admin/integrations/keys", {"name": "Teste", "scopes": ["orders:read"]}).json["id"]
    other = app.test_client()
    assert register(other, "outra-loja", "other@example.com").status_code == 201
    assert send(other, "DELETE", f"/api/admin/integrations/keys/{unused}/permanent", {}).status_code == 404
    assert send(client, "DELETE", f"/api/admin/integrations/keys/{unused}/permanent", {}).status_code == 200


def test_option_groups_required_sizes_and_half_and_half(shop):
    client, _ = shop
    category = client.get("/api/admin/products").json["categories"][0]["id"]
    pizza = {"name": "Pizza", "description": "", "price": 0, "category_id": category, "active": True, "unit_cost": 500,
             "extras": [{"id": "borda", "name": "Borda recheada", "price": 800, "unit_cost": 200}],
             "option_groups": [
                 {"id": "tam", "name": "Tamanho", "min": 1, "max": 1, "pricing": "sum",
                  "options": [{"id": "m", "name": "Média", "price": 3000, "unit_cost": 900}, {"id": "g", "name": "Grande", "price": 4000, "unit_cost": 1200}]},
                 {"id": "sab", "name": "Sabores", "min": 1, "max": 2, "pricing": "max",
                  "options": [{"id": "mu", "name": "Mussarela", "price": 0, "unit_cost": 300}, {"id": "ca", "name": "Calabresa", "price": 500, "unit_cost": 400}]}]}
    # Preço zero só é aceito quando há escolha obrigatória.
    assert send(client, "POST", "/api/admin/products", {**pizza, "option_groups": []}).status_code == 400
    bad = {**pizza, "option_groups": [{**pizza["option_groups"][0], "min": 3}]}
    assert send(client, "POST", "/api/admin/products", bad).status_code == 400
    pid = send(client, "POST", "/api/admin/products", pizza).json["id"]
    public = next(p for p in client.get("/api/store/minha-loja").json["products"] if p["id"] == pid)
    assert "unit_cost" not in public["option_groups"][0]["options"][0]

    def quote(ids):
        body = payload(pid, mode="pickup")
        body["items"][0]["option_ids"] = ids
        return send(client, "POST", "/api/store/minha-loja/quote", body)

    missing = quote(["mu"])
    assert missing.status_code == 400 and "Tamanho" in missing.json["error"]
    assert quote(["g", "mu", "ca", "m"]).status_code == 400
    # Grande (40) + meio a meio cobra o mais caro (5) + borda (8) = 53
    assert quote(["g", "mu", "ca", "borda"]).json["subtotal"] == 5300
    assert quote(["m", "mu"]).json["subtotal"] == 3000
    body = payload(pid, mode="pickup")
    body["items"][0]["option_ids"] = ["g", "mu", "ca"]
    assert send(client, "POST", "/api/store/minha-loja/orders", body, uuid.uuid4().hex).status_code == 201
    item = client.get("/api/admin/orders").json["orders"][0]["items"][0]
    assert item["unit_price"] == 4500 and item["unit_cost"] == 500 + 1200 + 400
    assert {x["name"] for x in item["extras"]} == {"Tamanho: Grande", "Sabores: Mussarela", "Sabores: Calabresa"}


def test_scheduled_orders_respect_hours_and_reviews(shop):
    from datetime import timedelta
    client, pid = shop
    store = client.get("/api/admin/store").json["store"]
    later = datetime.now(ZONE) + timedelta(days=2)
    day = str(later.weekday())
    # Loja fora do horário agora, aberta só no dia agendado das 10h às 22h.
    saved = send(client, "PUT", "/api/admin/store", {**store, "auto_hours": True, "hours": {day: [["10:00", "22:00"]]}})
    assert saved.status_code == 200
    good = later.replace(hour=19, minute=30).strftime("%Y-%m-%dT%H:%M")
    bad = later.replace(hour=23, minute=0).strftime("%Y-%m-%dT%H:%M")
    assert send(client, "POST", "/api/store/minha-loja/orders", payload(pid, scheduled_for=good), uuid.uuid4().hex).status_code == 400
    store = client.get("/api/admin/store").json["store"]
    assert send(client, "PUT", "/api/admin/store", {**store, "allow_scheduling": True}).status_code == 200
    assert send(client, "POST", "/api/store/minha-loja/orders", payload(pid, scheduled_for=bad), uuid.uuid4().hex).status_code == 400
    too_far = (datetime.now(ZONE) + timedelta(days=9)).strftime("%Y-%m-%dT%H:%M")
    assert send(client, "POST", "/api/store/minha-loja/orders", payload(pid, scheduled_for=too_far), uuid.uuid4().hex).status_code == 400
    if store_status({**store, "enabled": 1, "open": 1, "auto_hours": 1, "hours": '{"%s": [["10:00", "22:00"]]}' % day})[0] is False:
        assert send(client, "POST", "/api/store/minha-loja/orders", payload(pid), uuid.uuid4().hex).status_code == 409
    created = send(client, "POST", "/api/store/minha-loja/orders", payload(pid, scheduled_for=good), uuid.uuid4().hex)
    assert created.status_code == 201
    token = created.json["token"]
    tracked = client.get("/api/track/" + token).json["order"]
    assert tracked["scheduled_for"] and tracked["rating"] is None
    assert send(client, "POST", f"/api/track/{token}/review", {"rating": 5}).status_code == 409
    order = client.get("/api/admin/orders").json["orders"][0]
    advance(client, order["id"], "preparing", "ready", "delivering", "completed")
    assert send(client, "POST", f"/api/track/{token}/review", {"rating": 6}).status_code == 400
    assert send(client, "POST", f"/api/track/{token}/review", {"rating": 4, "comment": "Chegou quentinho"}).status_code == 200
    assert send(client, "POST", f"/api/track/{token}/review", {"rating": 1}).status_code == 409
    reviews = client.get("/api/admin/reviews").json
    assert reviews["average"] == 4.0 and reviews["reviews"][0]["review"] == "Chegou quentinho"


def test_team_roles_are_enforced_by_the_server(app, shop):
    client, pid = shop
    def member(email, role):
        assert send(client, "POST", "/api/admin/team", {"name": "Pessoa " + role, "email": email, "password": "senha-da-equipe-1", "role": role}).status_code == 201
        person = app.test_client()
        assert send(person, "POST", "/api/auth/login", {"email": email, "password": "senha-da-equipe-1"}).status_code == 200
        return person
    assert send(client, "POST", "/api/admin/team", {"name": "X", "email": "owner@example.com", "password": "senha-da-equipe-1", "role": "cashier"}).status_code in (400, 409)
    assert send(client, "POST", "/api/admin/team", {"name": "Xis", "email": "dono2@example.com", "password": "senha-da-equipe-1", "role": "owner"}).status_code == 400
    kitchen, cashier, manager = member("cozinha@example.com", "kitchen"), member("caixa@example.com", "cashier"), member("gerente@example.com", "manager")
    assert kitchen.get("/api/admin/store").json["role"] == "kitchen"
    order = make_order(client, pid)
    # Cozinha: vê e avança pedidos, mas não confirma pagamento, não cancela e não vê vendas.
    assert send(kitchen, "PATCH", f"/api/admin/orders/{order['id']}", {"status": "preparing", "paid": True}).status_code == 200
    assert not client.get("/api/admin/orders").json["orders"][0]["paid"]
    assert send(kitchen, "PATCH", f"/api/admin/orders/{order['id']}", {"status": "cancelled"}).status_code == 403
    for url in ("/api/admin/summary?days=1", "/api/admin/finance", "/api/admin/team", "/api/admin/report?days=7"):
        assert kitchen.get(url).status_code == 403
    assert send(kitchen, "PUT", "/api/admin/store", client.get("/api/admin/store").json["store"]).status_code == 403
    # Caixa: pedidos, balcão e vendas; sem cardápio, financeiro ou integrações.
    assert cashier.get("/api/admin/summary?days=1").status_code == 200
    assert send(cashier, "POST", "/api/admin/orders", payload(pid), uuid.uuid4().hex).status_code == 201
    assert send(cashier, "POST", "/api/admin/products", {"name": "Hack", "price": 1}).status_code == 403
    assert cashier.get("/api/admin/integrations").status_code == 403
    assert send(cashier, "POST", "/api/admin/store/open", {"open": False}).status_code == 200
    # Gerente: quase tudo, menos equipe e integrações.
    assert manager.get("/api/admin/finance").status_code == 200
    assert manager.get("/api/admin/team").status_code == 403
    assert send(manager, "POST", "/api/admin/integrations/keys", {"name": "x", "scopes": ["orders:read"]}).status_code == 403
    # Desativar derruba a sessão na hora; trocar a função também.
    members = client.get("/api/admin/team").json["members"]
    cid = next(m["id"] for m in members if m["email"] == "caixa@example.com")
    assert send(client, "PATCH", f"/api/admin/team/{cid}", {"active": False}).status_code == 200
    assert cashier.get("/api/admin/orders").status_code == 401
    kid = next(m["id"] for m in members if m["email"] == "cozinha@example.com")
    assert send(client, "PATCH", f"/api/admin/team/{kid}", {"role": "cashier"}).status_code == 200
    assert kitchen.get("/api/admin/orders").status_code == 401
    # Outra loja não enxerga nem altera a equipe desta.
    other = app.test_client()
    assert register(other, "outra-loja", "other@example.com").status_code == 201
    assert other.get("/api/admin/team").json["members"] == []
    assert send(other, "PATCH", f"/api/admin/team/{kid}", {"active": False}).status_code == 404


def test_push_subscription_and_new_order_alert(app, shop):
    client, pid = shop
    sent = []

    def fake_deliver(config, subscriptions, message):
        sent.append((len(subscriptions), message))
        return [s["endpoint"] for s in subscriptions if "morto" in s["endpoint"]]

    app.extensions["push_deliver"] = fake_deliver
    worker = client.get("/sw.js")
    assert worker.status_code == 200 and worker.headers["Service-Worker-Allowed"] == "/"
    key = client.get("/api/admin/push/key").json["public_key"]
    assert len(key) == 87 and client.get("/api/admin/push/key").json["public_key"] == key
    sub = {"endpoint": "https://push.example.com/aparelho-1", "keys": {"p256dh": "B" * 87, "auth": "a" * 22}}
    assert send(client, "POST", "/api/admin/push/subscribe", {**sub, "endpoint": "http://inseguro.example.com/x"}).status_code == 400
    assert send(client, "POST", "/api/admin/push/subscribe", sub).status_code == 201
    assert send(client, "POST", "/api/admin/push/subscribe", {**sub, "endpoint": "https://push.example.com/morto"}).status_code == 201
    make_order(client, pid)
    count, message = sent[-1]
    assert count == 2 and message["title"].startswith("Novo pedido #") and "Cliente Teste" in message["body"]
    # O aparelho que não existe mais é removido sozinho.
    assert send(client, "POST", "/api/admin/push/test", {}).json["sent"] == 1
    # Pedido de balcão não dispara aviso para a própria loja.
    before = len(sent)
    assert send(client, "POST", "/api/admin/orders", payload(pid), uuid.uuid4().hex).status_code == 201
    assert len(sent) == before
    assert send(client, "POST", "/api/admin/push/unsubscribe", {"endpoint": sub["endpoint"]}).status_code == 200
    assert send(client, "POST", "/api/admin/push/test", {}).json["sent"] == 0
